#!/usr/bin/env python3
"""table14_runtime.py — 打印 Table 14（tab:runtime）

R1 把原来的两张运行时表并进**同一个 table* 浮动体**，一个 \\label 下挂两个
tabular：

  (a) Base-scale, multi-GPU     Cases 43-44
      COMSOL(CPU) 与 1/2/4 卡 A800 的单样本时延、批量吞吐、相对加速比
  (b) Scaling with domain size  Cases 45-50
      单 DCU 上域尺寸 128→512 m 的单样本推理时间

两半口径不同，不可混比：
  · (a) 的吞吐是 200 样本并行的**批量吞吐**（总样本/壁钟）
  · (b) 只呈现**单样本时延**，取自训练结束(200 轮)最后一个"推理时间统计摘要"
    块；其吞吐（=1000/单样本ms，逐样本串行速率）只记在 xlsx，不进论文

取数一律复用核验脚本 ch4_validation/scripts/T14_runtime.py 的 load_base() /
load_scale()，与 verify.py 走同一次调用，两边不会各写一套解析而漂移。

    python table14_runtime.py [--tex]
"""
import _tblcommon as K


def table14_a():
    K.head("tab:runtime", "Cases 43-44 · 单轮计时与相对 COMSOL 加速比")
    xl = K.paths.xlsx_path("4.8")
    K.note(f"xlsx: {K.paths.rel(xl)}  (sheet 1)")
    K.note("取数复用 ch4_validation/scripts/T14_runtime.py 的 load_base()")
    xd = K.checker_module("T14_runtime").load_base()
    w = [10, 9, 11, 12, 11]
    K.row(["Case", "Method", "Time(ms)", "Thr.(samp/s)", "Speed-up"], w)
    K.rule(w)
    for no, tag in ((43, "R1"), (44, "W1")):
        d = xd.get(no, {})
        for m in ("COMSOL", "1 GPU", "2 GPU", "4 GPU"):
            b = d.get(m) or {}
            K.row([f"{no} ({tag})" if m == "COMSOL" else "",
                   m, K.f2(b.get("time")), K.f2(b.get("thr")),
                   K.f2(b.get("speedup"))], w)
        K.rule(w)
    if K.want_tex():
        print("\n  tex 数据行 (a)（浮动体内第一张 tabular）：")
        for r in K.tex_rows_of(0, "tab:runtime"):
            print("   ", " | ".join(r))


def table14_b():
    K.head("tab:runtime", "Cases 45-50 · 单 DCU 域尺度缩放（论文只列 Time）")
    xl = K.paths.xlsx_path("4.8")
    K.note(f"xlsx: {K.paths.rel(xl)}  (sheet 2)")
    K.note("Time 取自各案例最后一个推理时间统计摘要块（训练结束，200 轮）")
    K.note("取数复用 T14_runtime.py 的 load_scale()；吞吐另读 H 列")
    xd = K.checker_module("T14_runtime").load_scale()
    thr = _thr_col(xl)
    w = [6, 8, 7, 11, 11, 13]
    K.row(["Case", "Dataset", "Lx(m)", "N(nodes)", "Time(ms)",
           "Thr.(samp/s)*"], w)
    K.rule(w)
    for no in sorted(xd):
        d = xd[no]
        K.row([no, d["dataset"], d["lx"], f"{d['n']:,}",
               K.f2(d["time"]), K.f2(thr.get(no))], w)
    K.rule(w)
    K.note("* 吞吐 = 1000/Time（逐样本串行速率），仅 xlsx 记录，不进论文表/图")
    if K.want_tex():
        print("\n  tex 数据行 (b)（浮动体内第二张 tabular）：")
        for r in K.tex_rows_of(1, "tab:runtime"):
            print("   ", " | ".join(r))


def _thr_col(xl):
    """sheet 2 的吞吐量列（H）。论文不用，仅在此打印以便人工核对 xlsx。"""
    import openpyxl
    ws = openpyxl.load_workbook(xl, data_only=True).worksheets[1]
    out = {}
    for rr in range(5, ws.max_row + 1):
        v = ws.cell(rr, 1).value
        if v is None:
            continue
        h = ws.cell(rr, 8).value
        out[int(v)] = float(h) if h is not None else None
    return out


if __name__ == "__main__":
    table14_a()
    table14_b()
    print()
