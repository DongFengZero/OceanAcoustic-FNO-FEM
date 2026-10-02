#!/usr/bin/env python3
"""table20_21_runtime.py — 打印 Table 20（tab:runtime）与 Table 21（tab:runtime-scale）

Table 20  Cases 43-44：COMSOL(CPU) 与 1/2/4 卡 A800 的单样本时延、批量吞吐、加速比
Table 21  Cases 45-50：单 DCU 上域尺寸 128→512 m 的单样本推理时间

两表口径不同，不可混比：
  · Table 20 的吞吐是 200 样本并行的**批量吞吐**（总样本/壁钟）
  · Table 21 只呈现**单样本时延**，取自训练结束(200 轮)最后一个"推理时间统计摘要"块
    其吞吐（=1000/单样本ms，逐样本串行速率）只记在 xlsx，不进论文

    python table20_21_runtime.py [--tex]
"""
import _tblcommon as K


def table20():
    K.head("tab:runtime", "Cases 43-44 · 单轮计时与相对 COMSOL 加速比")
    xl = K.paths.xlsx_path("4.8")
    K.note(f"xlsx: {K.paths.rel(xl)}  (sheet 1)")
    K.note("取数函数复用 ch4_validation/scripts/T20_runtime.py 的 load_xlsx()")
    xd = K.checker_module("T20_runtime").load_xlsx()
    w = [10, 9, 11, 12, 11]
    K.row(["Case", "Method", "Time(ms)", "Thr.(samp/s)", "Speed-up"], w)
    K.rule(w)
    for no, tag in ((43, "R1"), (44, "W1")):
        d = xd.get(no, {})
        for m in ("COMSOL", "1 GPU", "2 GPU", "4 GPU"):
            b = d.get(m) or {}
            K.row([f"{no} ({tag})" if m == "COMSOL" else "",
                   m, K.f2(b.get("time")), K.f2(b.get("thr")),
                   K.f2(b.get("speedup"))], w)
        K.rule(w)
    if K.want_tex():
        print("\n  tex 数据行：")
        for r in K.tex_rows("tab:runtime"):
            print("   ", " | ".join(r))


def table21():
    K.head("tab:runtime-scale", "Cases 45-50 · 单 DCU 域尺度缩放（论文只列 Time）")
    xl = K.paths.xlsx_path("4.8")
    K.note(f"xlsx: {K.paths.rel(xl)}  (sheet 2)")
    K.note("Time 取自各案例最后一个推理时间统计摘要块（训练结束，200 轮）")
    K.note("Case/Lx/N/Time 复用 T21_runtime_scale.py 的 load_xlsx()；吞吐另读 H 列")
    xd = K.checker_module("T21_runtime_scale").load_xlsx()
    thr = _thr_col(xl)
    w = [6, 8, 7, 11, 11, 13]
    K.row(["Case", "Dataset", "Lx(m)", "N(nodes)", "Time(ms)",
           "Thr.(samp/s)*"], w)
    K.rule(w)
    for no in sorted(xd):
        d = xd[no]
        K.row([no, d["dataset"], d["lx"], f"{d['n']:,}",
               K.f2(d["time"]), K.f2(thr.get(no))], w)
    K.rule(w)
    K.note("* 吞吐 = 1000/Time（逐样本串行速率），仅 xlsx 记录，不进论文表/图")
    if K.want_tex():
        print("\n  tex 数据行：")
        for r in K.tex_rows("tab:runtime-scale"):
            print("   ", " | ".join(r))


def _thr_col(xl):
    """sheet 2 的吞吐量列（H）。论文不用，仅在此打印以便人工核对 xlsx。"""
    import openpyxl
    ws = openpyxl.load_workbook(xl, data_only=True).worksheets[1]
    out = {}
    for rr in range(5, ws.max_row + 1):
        v = ws.cell(rr, 1).value
        if v is None:
            continue
        h = ws.cell(rr, 8).value
        out[int(v)] = float(h) if h is not None else None
    return out


if __name__ == "__main__":
    table20()
    table21()
    print()
